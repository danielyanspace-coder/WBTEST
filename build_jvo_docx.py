# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

ACCENT = RGBColor(0x4A, 0x1F, 0x8C)   # фиолетовый — фирменный для AI
DARK   = RGBColor(0x1F, 0x1F, 0x1F)
GREY   = RGBColor(0x77, 0x77, 0x77)

def heading(text, size=15, color=ACCENT, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(' — ' + text)
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.size = Pt(10.5)
    return p

def para(text, size=11, color=DARK, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.italic = italic
    return p

# ===== Заголовок сверху: КарЕг =====
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('КарЕг')
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = ACCENT

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run('Полный разбор сервиса JVO (jvo.ru)')
rs.font.size = Pt(13)
rs.font.color.rgb = GREY
rs.italic = True

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
rm = meta.add_run('Аналитика и AI-автоматизация для Wildberries и Ozon  •  24.06.2026')
rm.font.size = Pt(9)
rm.font.color.rgb = GREY

# разделитель
hr = doc.add_paragraph()
pPr = hr._p.get_or_add_pPr()
pbdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '4A1F8C')
pbdr.append(bottom); pPr.append(pbdr)

# ===== 1. О сервисе =====
heading('1. О сервисе', space_before=6)
para('JVO — российская платформа аналитики и AI-автоматизации для продавцов на '
     'маркетплейсах, запущенная в 2022 году. Позиционируется как «персональный AI-менеджер», '
     'который управляет ценами, отвечает на отзывы, следит за остатками и анализирует данные '
     '24/7. Сервис делает ставку на современный UX и встроенный искусственный интеллект там, '
     'где «старые» платформы ограничиваются классическими таблицами.')
bullet('поддерживает одновременно Wildberries и Ozon (оба в одной подписке)', 'Площадки')
bullet('аккредитована как IT-компания, технологический партнёр Ozon, авторизованный сервис Wildberries', 'Статус')
bullet('собственная Big Data: учёт поведения покупателей и динамики рынка', 'Данные')
bullet('экономия до 500 часов ручной работы в месяц, до 80–90% операционных задач — на AI', 'Эффект')

# ===== 2. Архитектура: AI-агенты и модули =====
heading('2. Ключевые компоненты платформы')
bullet('первый в мире AI-агент для e-commerce; забирает до 90% ручных задач', 'JVO Агент')
bullet('автоматическое управление ценами по 20+ параметрам', 'Агент ценообразования')
bullet('автоответы на отзывы и вопросы в Tone of Voice бренда', 'Агент коммуникаций')
bullet('автоматизация рекламных кампаний', 'Агент рекламы')
bullet('SEO-оптимизация карточек товаров', 'SEO Pro')
bullet('единый мониторинг ключевых показателей бизнеса', 'Дашборд')
bullet('управление логистикой и остатками', 'Логистика')
bullet('система событий и уведомлений', 'События')

# ===== 3. Управление ценами =====
heading('3. Умное ценообразование (Агент ценообразования)')
for x in [
    'Анализ 20+ параметров: спрос, остатки, оборачиваемость, рейтинг, отзывы, конверсия в корзину',
    'Изменение цен по расписанию — с подтверждением или в полном автомате',
    'Удержание маржинальности при разгоне продаж',
    'Предотвращение Out-of-stock (дефицита товара)',
    'Стимулирование продаж через динамическое ценообразование',
]:
    bullet(x)

# ===== 4. Коммуникации и отзывы =====
heading('4. Работа с отзывами и коммуникациями (Агент коммуникаций)')
for x in [
    'Автоответы на отзывы покупателей',
    'Автоответы на вопросы покупателей',
    'Ответы в Tone of Voice бренда (без шаблонных фраз)',
    'Ответы в заданное время / по расписанию',
    'Специальная обработка негатива для снижения ущерба рейтингу',
    'Кросс-продажи: рекомендация товаров в ответах на позитивные отзывы',
    'Анализ отзывов: выделение повторяющихся проблем и частоты упоминания недостатков',
    'Отчёты с конкретными действиями для улучшения товара',
]:
    bullet(x)

# ===== 5. Аналитика =====
heading('5. Аналитика и мониторинг')
for x in [
    'Аналитика Wildberries и Ozon в одном кабинете',
    'Аналитика товарной матрицы',
    'Анализ ключевых факторов ранжирования',
    'Визуализация данных (понятные дашборды вместо «портянки» из 50 колонок)',
    'Big Data-анализ на собственной базе с учётом поведения покупателей',
    'Предупреждения о критичных ошибках',
    'Оповещения об out-of-stock (полном или частичном)',
    'Автоматическое формирование задач для предотвращения потерь дохода',
]:
    bullet(x)

# ===== 6. Реклама и SEO =====
heading('6. Реклама и SEO')
for x in [
    'Агент рекламы — автоматизация и оптимизация рекламных кампаний',
    'SEO Pro — оптимизация карточек товаров под поиск',
]:
    bullet(x)

# ===== 7. Тарифы и аудитория =====
heading('7. Тарифы и целевая аудитория')
para('Оба маркетплейса входят в подписку. Конкретные цены публикуются в разделе «Тарифы» '
     'на сайте и зависят от оборота. Сервис ориентирован на селлеров с оборотом примерно '
     'от 1 млн до 2+ млн ₽/мес, а также на тех, кто приходит в аналитику без бухгалтерского '
     'бэкграунда и хочет «понятные» дашборды.', color=DARK)

# ===== 8. Кейсы =====
heading('8. Результаты и кейсы')
for x in [
    'Орехи/сухофрукты: +34 млн ₽ за месяц',
    'Косметика: +10 млн ₽ за месяц',
    'Одежда: прирост 26% в первый месяц',
    'Agroimpex: выручка на WB ×2,7 — до 306 млн ₽ за год без расширения штата',
    'Первые клиенты AI-агента: рост конверсии на 12–18% за первый месяц при падении среднего чека всего на 3–5%',
]:
    bullet(x)

# ===== 9. Итог =====
heading('9. Сильные стороны (резюме)')
for x in [
    'Глубокая AI-автоматизация рутины (цены, отзывы, реклама)',
    'WB + Ozon в одной подписке',
    'Современный UX и наглядная визуализация',
    'Официальная аккредитация и партнёрство с площадками',
    'Подходит новичкам без аналитического бэкграунда',
]:
    bullet(x)

# Подвал
doc.add_paragraph()
note = doc.add_paragraph()
rn = note.add_run('Источники: jvo.ru (разделы Агент, Коммуникации, Аналитика), MPAgency, '
                  'Uniseller, Forbes.ru, vc.ru. Данные собраны из открытых обзоров и страниц сервиса.')
rn.font.size = Pt(8)
rn.font.color.rgb = GREY

out = '/home/user/WBTEST/JVO_разбор_сервиса.docx'
doc.save(out)
print('Saved:', out)
